"""
Módulo de geração de relatórios de ocupação de salas em formato PDF e Excel (XLSX).
Design com fallback resiliente para execução garantida em ambientes desacoplados.
"""

import io
import datetime
from typing import List, Dict, Any, Optional

try:
    from reportlab.lib.pagesizes import letter
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib import colors
    HAS_REPORTLAB = True
except ImportError:
    HAS_REPORTLAB = False

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False


def filter_data(rooms: List[Dict[str, Any]], allocations: List[Dict[str, Any]], block_id: Optional[str] = None, shift: Optional[str] = None):
    """
    Filtra salas e alocações com base nos parâmetros block_id e shift.
    """
    filtered_rooms = rooms
    if block_id and block_id != "all":
        filtered_rooms = [r for r in rooms if r.get("block_id") == block_id]

    shift_slots = None
    if shift and shift != "all":
        s_upper = shift.upper()
        if s_upper == 'M':
            shift_slots = {'M1', 'M2', 'M3', 'M4', 'M5', 'M6'}
        elif s_upper == 'T':
            shift_slots = {'T1', 'T2', 'T3', 'T4', 'T5', 'T6'}
        elif s_upper == 'N':
            shift_slots = {'N1', 'N2', 'N3', 'N4', 'N5', 'N6'}

    filtered_allocations = allocations
    if shift_slots:
        filtered_allocations = [a for a in allocations if a.get("time_slot") in shift_slots or (a.get("shift") and a.get("shift").upper() == shift.upper())]

    if block_id and block_id != "all":
        room_ids = {r.get("id") or r.get("name") for r in filtered_rooms}
        filtered_allocations = [a for a in filtered_allocations if a.get("room_id") in room_ids]

    return filtered_rooms, filtered_allocations


def compute_summary_by_block(rooms: List[Dict[str, Any]], allocations: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    """
    Consolida as estatísticas de ocupação agrupadas por bloco predial.
    """
    blocks = {}
    for r in rooms:
        b_id = r.get("block_id", "Geral")
        if b_id not in blocks:
            blocks[b_id] = {
                "block_id": b_id,
                "total_rooms": 0,
                "total_capacity": 0,
                "occupied_slots": 0,
                "room_ids": set()
            }
        blocks[b_id]["total_rooms"] += 1
        blocks[b_id]["total_capacity"] += r.get("capacity", 0)
        r_key = r.get("id") or r.get("name")
        if r_key:
            blocks[b_id]["room_ids"].add(r_key)

    for alloc in allocations:
        r_id = alloc.get("room_id")
        for b_info in blocks.values():
            if r_id in b_info["room_ids"]:
                b_info["occupied_slots"] += 1

    summary = []
    for b_id, b_info in sorted(blocks.items()):
        total_available = b_info["total_rooms"] * 6  # 6 slots padrão por turno
        occupied = b_info["occupied_slots"]
        rate = round((occupied / total_available * 100), 1) if total_available > 0 else 0.0
        summary.append({
            "block_id": b_id,
            "total_rooms": b_info["total_rooms"],
            "total_capacity": b_info["total_capacity"],
            "occupied_slots": occupied,
            "total_available_slots": total_available,
            "occupancy_rate": rate
        })

    return summary


def compute_occupancy_matrix(rooms: List[Dict[str, Any]], allocations: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    """
    Gera a matriz de ocupação de cada sala por slot.
    """
    matrix = []
    for r in rooms:
        r_id = r.get("id") or r.get("name")
        r_name = r.get("name", "Sala")
        b_id = r.get("block_id", "Geral")
        cap = r.get("capacity", 0)
        
        r_allocs = [a for a in allocations if a.get("room_id") == r_id or a.get("room_id") == r_name]
        occupied_count = len(r_allocs)
        rate = round((occupied_count / 18.0) * 100, 1) if 18 > 0 else 0.0
        
        matrix.append({
            "room_id": r_id,
            "room_name": r_name,
            "block_id": b_id,
            "capacity": cap,
            "occupied_count": occupied_count,
            "occupancy_rate": rate,
            "allocations": r_allocs
        })
    return matrix


def generate_pdf_report(rooms: List[Dict[str, Any]], allocations: List[Dict[str, Any]], block_id: Optional[str] = None, shift: Optional[str] = None, teacher_allocations: Optional[List[Dict[str, Any]]] = None) -> bytes:
    """
    Gera relatório consolidado de ocupação geral em PDF.
    """
    filtered_rooms, filtered_allocs = filter_data(rooms, allocations, block_id, shift)
    summary = compute_summary_by_block(filtered_rooms, filtered_allocs)

    if not HAS_REPORTLAB:
        output = io.StringIO()
        output.write(f"CLASSSYNC AI - RELATÓRIO DE OCUPAÇÃO GERAL\n")
        output.write(f"Emissão: {datetime.datetime.now().strftime('%d/%m/%Y %H:%M')}\n")
        output.write("=" * 60 + "\n")
        for s in summary:
            output.write(f"Bloco: {s['block_id']} | Salas: {s['total_rooms']} | Cap: {s['total_capacity']} | Ocup: {s['occupied_slots']}/{s['total_available_slots']} ({s['occupancy_rate']}%)\n")
        return output.getvalue().encode("utf-8")

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, rightMargin=36, leftMargin=36, topMargin=36, bottomMargin=36)
    styles = getSampleStyleSheet()

    elements = []
    title_style = ParagraphStyle('RepTitle', parent=styles['Heading1'], fontName='Helvetica-Bold', fontSize=18, textColor=colors.HexColor('#0f172a'), spaceAfter=6)
    sub_style = ParagraphStyle('RepSub', parent=styles['Normal'], fontName='Helvetica', fontSize=10, textColor=colors.HexColor('#475569'), spaceAfter=14)

    elements.append(Paragraph("ClassSync AI — Relatório de Ocupação Predial", title_style))
    elements.append(Paragraph(f"Filtros: Bloco={block_id or 'Todos'} | Turno={shift or 'Todos'} | Emissão: {datetime.datetime.now().strftime('%d/%m/%Y %H:%M')}", sub_style))
    elements.append(Spacer(1, 8))

    table_data = [["Bloco", "Qtd Salas", "Capacidade Total", "Slots Ocupados", "Taxa Ocupação %"]]
    for s in summary:
        table_data.append([
            s["block_id"],
            str(s["total_rooms"]),
            f"{s['total_capacity']} pessoas",
            f"{s['occupied_slots']} / {s['total_available_slots']}",
            f"{s['occupancy_rate']}%"
        ])

    t = Table(table_data, colWidths=[110, 85, 115, 115, 115])
    t.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#0f172a')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 9),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 8),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f8fafc')]),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#cbd5e1')),
        ('ALIGN', (1, 0), (-1, -1), 'CENTER'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('PADDING', (0, 0), (-1, -1), 6),
    ]))
    elements.append(t)
    doc.build(elements)

    pdf_bytes = buffer.getvalue()
    buffer.close()
    return pdf_bytes


def generate_excel_report(rooms: List[Dict[str, Any]], allocations: List[Dict[str, Any]], block_id: Optional[str] = None, shift: Optional[str] = None, teacher_allocations: Optional[List[Dict[str, Any]]] = None) -> bytes:
    """
    Gera relatório consolidado de ocupação geral em Excel (.xlsx).
    """
    filtered_rooms, filtered_allocs = filter_data(rooms, allocations, block_id, shift)
    summary = compute_summary_by_block(filtered_rooms, filtered_allocs)

    if not HAS_OPENPYXL:
        output = io.StringIO()
        output.write("Bloco;Salas;Capacidade;Ocupados;Disponiveis;Taxa\n")
        for s in summary:
            output.write(f"{s['block_id']};{s['total_rooms']};{s['total_capacity']};{s['occupied_slots']};{s['total_available_slots']};{s['occupancy_rate']}\n")
        return output.getvalue().encode("utf-8-sig")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Ocupação de Ambientes"

    headers = ["Bloco", "Salas", "Capacidade Total", "Slots Ocupados", "Slots Disponíveis", "Taxa Ocupação %"]
    ws.append(headers)

    for s in summary:
        ws.append([s["block_id"], s["total_rooms"], s["total_capacity"], s["occupied_slots"], s["total_available_slots"], s["occupancy_rate"]])

    buffer = io.BytesIO()
    wb.save(buffer)
    excel_bytes = buffer.getvalue()
    buffer.close()
    return excel_bytes


def _parse_day_num(d: Any) -> int:
    if isinstance(d, int):
        return d if 1 <= d <= 7 else 1
    d_str = str(d or "").strip().lower()
    if d_str.isdigit():
        val = int(d_str)
        return val if 1 <= val <= 7 else 1
    mapping = {
        "segunda": 1, "segunda-feira": 1, "seg": 1,
        "terça": 2, "terca": 2, "terça-feira": 2, "terca-feira": 2, "ter": 2,
        "quarta": 3, "quarta-feira": 3, "qua": 3,
        "quinta": 4, "quinta-feira": 4, "qui": 4,
        "sexta": 5, "sexta-feira": 5, "sex": 5,
        "sábado": 6, "sabado": 6, "sab": 6,
        "domingo": 7, "dom": 7
    }
    return mapping.get(d_str, 1)


def _parse_shift_code(s: Any, time_slot: str = "") -> str:
    s_str = str(s or "").upper().strip()
    if s_str in ("M", "MANHÃ", "MATUTINO"):
        return "M"
    if s_str in ("T", "TARDE", "VESPERTINO"):
        return "T"
    if s_str in ("N", "NOITE", "NOTURNO"):
        return "N"
    if time_slot and time_slot[0].upper() in ("M", "T", "N"):
        return time_slot[0].upper()
    return "M"


def _parse_sub_slot_num(sub: Any, time_slot: str = "") -> int:
    if isinstance(sub, int) and 1 <= sub <= 6:
        return sub
    sub_str = str(sub or "").strip()
    if sub_str.isdigit():
        val = int(sub_str)
        return val if 1 <= val <= 6 else 1
    if time_slot and len(time_slot) >= 2 and time_slot[1:].isdigit():
        val = int(time_slot[1:])
        return val if 1 <= val <= 6 else 1
    return 1


def generate_teacher_pdf_report(teacher_name: str, teacher_dept: str, allocations: List[Dict[str, Any]]) -> bytes:
    """
    Gera um relatório PDF executivo com a grade semanal de horários do docente:
    - Metadados do Docente (Nome, Departamento, Carga Horária)
    - Indicadores Rápidos (Total de Aulas, Disciplinas Distintas, Turnos Ocupados)
    - Matriz Semanal de Horários (Grid Semanal x Slots M1..N6)
    - Detalhamento Sequencial Ordenado por Dia da Semana, Turno e Slot.
    """
    days_map = {1: "Segunda-feira", 2: "Terça-feira", 3: "Quarta-feira", 4: "Quinta-feira", 5: "Sexta-feira", 6: "Sábado", 7: "Domingo"}
    shift_map = {"M": "Manhã", "T": "Tarde", "N": "Noite"}
    shift_order = {"M": 1, "T": 2, "N": 3}

    subslot_times_map = {
        "M": {1: "07:00-07:50", 2: "07:50-08:40", 3: "08:40-09:30", 4: "09:45-10:35", 5: "10:35-11:25", 6: "11:25-12:15"},
        "T": {1: "13:00-13:50", 2: "13:50-14:40", 3: "14:40-15:30", 4: "15:45-16:35", 5: "16:35-17:25", 6: "17:25-18:15"},
        "N": {1: "19:00-19:50", 2: "19:50-20:40", 3: "20:40-21:30", 4: "21:45-22:35", 5: "22:35-23:25", 6: "23:25-00:15"}
    }

    norm_allocs = []
    for a in (allocations or []):
        day_n = _parse_day_num(a.get("day_of_week", 1))
        t_slot = a.get("time_slot", "")
        sh_code = _parse_shift_code(a.get("shift"), t_slot)
        sub_n = _parse_sub_slot_num(a.get("sub_slot"), t_slot)
        slot_code = t_slot if t_slot else f"{sh_code}{sub_n}"
        t_interval = a.get("time_interval") or subslot_times_map.get(sh_code, {}).get(sub_n, "07:00-07:50")
        
        norm_allocs.append({
            "subject_name": a.get("subject_name") or a.get("subject") or "Disciplina",
            "room_name": a.get("room_name") or a.get("room_id") or "Sala",
            "block_id": a.get("block_id") or "Bloco Geral",
            "shift": sh_code,
            "sub_slot": sub_n,
            "time_slot": slot_code,
            "time_interval": t_interval,
            "day_of_week": day_n
        })

    # Ordenar rigorosamente por dia da semana (1-7), turno (M, T, N) e sub_slot (1-6)
    sorted_allocs = sorted(
        norm_allocs,
        key=lambda x: (
            int(x["day_of_week"]),
            shift_order.get(x["shift"], 1),
            int(x["sub_slot"])
        )
    )

    total_aulas = len(sorted_allocs)
    unique_subjects = len(set(a["subject_name"] for a in sorted_allocs)) if sorted_allocs else 0
    unique_shifts = sorted(list(set(a["shift"] for a in sorted_allocs)), key=lambda s: shift_order.get(s, 1))
    turnos_ocupados_str = ", ".join(shift_map.get(s, s) for s in unique_shifts) if unique_shifts else "Nenhum"
    carga_estimada = round(total_aulas * 0.83, 1)

    if not HAS_REPORTLAB:
        lines = [
            "=" * 65,
            "CLASSSYNC AI - GRADE INDIVIDUAL DO DOCENTE",
            "=" * 65,
            f"Docente: {teacher_name}",
            f"Departamento: {teacher_dept} | Carga Horária: ~{carga_estimada}h/semana",
            f"Emissão: {datetime.datetime.now().strftime('%d/%m/%Y %H:%M')}",
            "-" * 65,
            f"MÉTRICAS: Total de Aulas: {total_aulas} | Disciplinas: {unique_subjects} | Turnos: {turnos_ocupados_str}",
            "-" * 65,
            "DISCIPLINAS, SALAS E HORÁRIOS:",
        ]
        for a in sorted_allocs:
            s_name = shift_map.get(a['shift'], a['shift'])
            d_name = days_map.get(a['day_of_week'], f"Dia {a['day_of_week']}")
            sub_num = f"{a['sub_slot']:02d}"
            lines.append(f"- Disciplina: {a['subject_name']} | Sala: {a['room_name']} ({a['block_id']}) | Dia: {d_name} | Turno: {s_name} | Horário: {a['time_interval']} (Slot: {a['time_slot']} - Aula {sub_num})")
        return "\n".join(lines).encode("utf-8")

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        rightMargin=30,
        leftMargin=30,
        topMargin=30,
        bottomMargin=30
    )

    styles = getSampleStyleSheet()
    
    app_header_style = ParagraphStyle(
        'TeacherAppHeader',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=9.5,
        textColor=colors.HexColor('#6366f1'),
        spaceAfter=3,
        textTransform='uppercase'
    )
    
    title_style = ParagraphStyle(
        'TeacherMainTitle',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=15,
        textColor=colors.HexColor('#0f172a'),
        spaceAfter=4
    )
    
    meta_style = ParagraphStyle(
        'TeacherMetaSub',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=8.5,
        textColor=colors.HexColor('#475569'),
        spaceAfter=8
    )

    section_heading = ParagraphStyle(
        'SectionHeading',
        parent=styles['Heading2'],
        fontName='Helvetica-Bold',
        fontSize=10.5,
        textColor=colors.HexColor('#1e293b'),
        spaceBefore=8,
        spaceAfter=4
    )

    kpi_label_style = ParagraphStyle(
        'KPILabel',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=7.5,
        textColor=colors.HexColor('#64748b'),
        textTransform='uppercase'
    )
    kpi_val_blue = ParagraphStyle(
        'KPIValBlue',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=10.5,
        textColor=colors.HexColor('#0284c7')
    )
    kpi_val_green = ParagraphStyle(
        'KPIValGreen',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=10.5,
        textColor=colors.HexColor('#16a34a')
    )
    kpi_val_purple = ParagraphStyle(
        'KPIValPurple',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=10.5,
        textColor=colors.HexColor('#9333ea')
    )

    grid_header_style = ParagraphStyle(
        'GridHeader',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=8,
        textColor=colors.white,
        alignment=1
    )

    grid_slot_cell = ParagraphStyle(
        'GridSlotCell',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=7.5,
        textColor=colors.HexColor('#0f172a'),
        alignment=1
    )

    grid_cell_text = ParagraphStyle(
        'GridCellText',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=7,
        leading=9,
        alignment=0
    )

    footer_style = ParagraphStyle(
        'DocFooter',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=7.5,
        textColor=colors.HexColor('#94a3b8'),
        alignment=1,
        spaceBefore=10
    )

    elements = []

    # 1. Cabeçalho Principal
    elements.append(Paragraph("ClassSync AI — Gestão de Espaços e Alocação", app_header_style))
    elements.append(Paragraph(f"Grade Semanal de Horários: {teacher_name}", title_style))
    elements.append(Paragraph(
        f"<b>Departamento:</b> {teacher_dept} &nbsp;|&nbsp; <b>Carga Estimada:</b> ~{carga_estimada}h ({total_aulas} aulas) &nbsp;|&nbsp; <b>Emissão:</b> {datetime.datetime.now().strftime('%d/%m/%Y às %H:%M')}",
        meta_style
    ))
    elements.append(Spacer(1, 2))

    # 2. Quadro de Métricas / Resumo Rápido (KPI Cards)
    kpi_col1 = [
        Paragraph("TOTAL DE AULAS", kpi_label_style),
        Paragraph(f"{total_aulas} {'aula alocada' if total_aulas == 1 else 'aulas alocadas'}", kpi_val_blue)
    ]
    kpi_col2 = [
        Paragraph("DISCIPLINAS DISTINTAS", kpi_label_style),
        Paragraph(f"{unique_subjects}", kpi_val_green)
    ]
    kpi_col3 = [
        Paragraph("TURNOS OCUPADOS", kpi_label_style),
        Paragraph(turnos_ocupados_str, kpi_val_purple)
    ]

    kpi_table = Table([[kpi_col1, kpi_col2, kpi_col3]], colWidths=[184, 184, 184])
    kpi_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#f8fafc')),
        ('BOX', (0, 0), (-1, -1), 1, colors.HexColor('#e2e8f0')),
        ('INNERGRID', (0, 0), (-1, -1), 1, colors.HexColor('#e2e8f0')),
        ('PADDING', (0, 0), (-1, -1), 6),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
    ]))
    elements.append(kpi_table)
    elements.append(Spacer(1, 6))

    # 3. Matriz Semanal de Horários (Grid Matriz)
    elements.append(Paragraph("Matriz Semanal de Horários (Visão Grade)", section_heading))
    
    all_possible_slots = [f"{sh}{num}" for sh in ["M", "T", "N"] for num in range(1, 7)]
    active_slots = [s for s in all_possible_slots if any(a["time_slot"] == s for a in sorted_allocs)]
    if not active_slots and sorted_allocs:
        active_slots = list(set(a["time_slot"] for a in sorted_allocs))

    grid_table_data = [[
        Paragraph("<b>Slot / Dia</b>", grid_header_style),
        Paragraph("<b>Segunda</b>", grid_header_style),
        Paragraph("<b>Terça</b>", grid_header_style),
        Paragraph("<b>Quarta</b>", grid_header_style),
        Paragraph("<b>Quinta</b>", grid_header_style),
        Paragraph("<b>Sexta</b>", grid_header_style),
        Paragraph("<b>Sábado</b>", grid_header_style),
    ]]

    if active_slots:
        for slot in active_slots:
            sh_code = slot[0].upper() if len(slot) > 0 else "M"
            sub_n = int(slot[1:]) if len(slot) > 1 and slot[1:].isdigit() else 1
            t_interval = subslot_times_map.get(sh_code, {}).get(sub_n, "")
            
            row = [Paragraph(f"<b>{slot}</b><br/><font size=6 color='#64748b'>{t_interval}</font>", grid_slot_cell)]
            
            for day_idx in range(1, 7):
                matched = [a for a in sorted_allocs if a["time_slot"] == slot and a["day_of_week"] == day_idx]
                if matched:
                    cell_html_parts = []
                    for m in matched:
                        cell_html_parts.append(f"<b><font color='#0f172a'>{m['subject_name']}</font></b><br/><font color='#059669'>📍 {m['room_name']}</font>")
                    row.append(Paragraph("<br/>".join(cell_html_parts), grid_cell_text))
                else:
                    row.append(Paragraph("<font color='#cbd5e1'>-</font>", ParagraphStyle('CenterDash', parent=grid_cell_text, alignment=1)))
            grid_table_data.append(row)
    else:
        grid_table_data.append([
            Paragraph("<i>Nenhuma aula alocada nesta grade.</i>", styles['Normal']),
            "-", "-", "-", "-", "-", "-"
        ])

    # Total width 552: 72 + 80 * 6 = 552
    grid_table = Table(grid_table_data, colWidths=[72, 80, 80, 80, 80, 80, 80])
    grid_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#0f172a')),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 5),
        ('TOPPADDING', (0, 0), (-1, 0), 5),
        ('BACKGROUND', (0, 1), (0, -1), colors.HexColor('#f1f5f9')),
        ('ROWBACKGROUNDS', (1, 1), (-1, -1), [colors.white, colors.HexColor('#f8fafc')]),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#cbd5e1')),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('PADDING', (0, 0), (-1, -1), 4),
    ]))
    elements.append(grid_table)
    elements.append(Spacer(1, 8))

    # 4. Tabela de Detalhamento Sequencial
    elements.append(Paragraph("Detalhamento da Carga Horária Docente", section_heading))
    
    table_data = [[
        Paragraph("<font color='white'><b>Disciplina</b></font>", styles['Normal']),
        Paragraph("<font color='white'><b>Sala Alocada</b></font>", styles['Normal']),
        Paragraph("<font color='white'><b>Dia da Semana</b></font>", styles['Normal']),
        Paragraph("<font color='white'><b>Turno</b></font>", styles['Normal']),
        Paragraph("<font color='white'><b>Horário / Slot</b></font>", styles['Normal'])
    ]]

    shift_colors_map = {
        "Manhã": ('#eff6ff', '#1d4ed8'),
        "Tarde": ('#fffbeb', '#b45309'),
        "Noite": ('#faf5ff', '#7e22ce')
    }

    if sorted_allocs:
        for a in sorted_allocs:
            subj_name = a['subject_name']
            disciplina_cell = f"<b>{subj_name}</b>"

            r_name = a['room_name']
            b_id = a['block_id']
            sala_cell = f"<b>{r_name}</b>"
            if b_id:
                sala_cell += f"<br/><font size=7.5 color='#64748b'>{b_id}</font>"

            day_name = days_map.get(a['day_of_week'], f"Dia {a['day_of_week']}")

            shift_code = a['shift']
            shift_name = shift_map.get(shift_code, shift_code)
            shift_bg, shift_fg = shift_colors_map.get(shift_name, ('#eff6ff', '#1d4ed8'))
            turno_cell = f"<font color='{shift_fg}'><b>{shift_name}</b></font>"

            time_interval = a['time_interval']
            sub_num = f"{a['sub_slot']:02d}"
            slot_code = a['time_slot']
            horario_cell = f"<b>{time_interval}</b><br/><font size=7.5 color='#64748b'>Slot: {slot_code} (Aula {sub_num})</font>"

            table_data.append([
                Paragraph(disciplina_cell, styles['Normal']),
                Paragraph(sala_cell, styles['Normal']),
                Paragraph(day_name, styles['Normal']),
                Paragraph(turno_cell, styles['Normal']),
                Paragraph(horario_cell, styles['Normal'])
            ])
    else:
        table_data.append([
            Paragraph("<i>Nenhuma aula alocada para este docente no momento.</i>", styles['Normal']),
            "-", "-", "-", "-"
        ])

    data_table = Table(table_data, colWidths=[140, 135, 95, 65, 117])
    data_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#0f172a')),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 6),
        ('TOPPADDING', (0, 0), (-1, 0), 6),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f8fafc')]),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#cbd5e1')),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('PADDING', (0, 0), (-1, -1), 5),
    ]))
    elements.append(data_table)

    # 5. Rodapé institucional
    elements.append(Spacer(1, 10))
    elements.append(Paragraph(
        "Documento gerado automaticamente pelo <b>ClassSync AI</b> • Relatório Oficial da Grade Docente",
        footer_style
    ))

    doc.build(elements)
    pdf_bytes = buffer.getvalue()
    buffer.close()
    return pdf_bytes


def generate_room_pdf_report(
    room_name: str, 
    block_id: str, 
    capacity: int, 
    room_type: str, 
    is_accessible: bool, 
    allocations: List[Dict[str, Any]]
) -> bytes:
    """
    Gera um relatório PDF fiel e completo com todas as informações do modal de ocupação:
    - Metadados da Sala (Nome, Bloco, Tipo, Capacidade, Acessibilidade)
    - Indicadores Rápidos (Total de Aulas, Docentes Distintos, Turnos Ocupados)
    - Tabela Estruturada (Docente, Disciplina, Dia da Semana, Turno, Horário / Slot)
    """
    days_map = {1: "Segunda-feira", 2: "Terça-feira", 3: "Quarta-feira", 4: "Quinta-feira", 5: "Sexta-feira", 6: "Sábado", 7: "Domingo"}
    shift_map = {"M": "Manhã", "T": "Tarde", "N": "Noite"}
    shift_order = {"M": 1, "T": 2, "N": 3}

    subslot_times_map = {
        "M": {1: "07:00-07:50", 2: "07:50-08:40", 3: "08:40-09:30", 4: "09:45-10:35", 5: "10:35-11:25", 6: "11:25-12:15"},
        "T": {1: "13:00-13:50", 2: "13:50-14:40", 3: "14:40-15:30", 4: "15:45-16:35", 5: "16:35-17:25", 6: "17:25-18:15"},
        "N": {1: "19:00-19:50", 2: "19:50-20:40", 3: "20:40-21:30", 4: "21:45-22:35", 5: "22:35-23:25", 6: "23:25-00:15"}
    }

    norm_allocs = []
    for a in (allocations or []):
        day_n = _parse_day_num(a.get("day_of_week", 1))
        t_slot = a.get("time_slot", "")
        sh_code = _parse_shift_code(a.get("shift"), t_slot)
        sub_n = _parse_sub_slot_num(a.get("sub_slot"), t_slot)
        slot_code = t_slot if t_slot else f"{sh_code}{sub_n}"
        t_interval = a.get("time_interval") or subslot_times_map.get(sh_code, {}).get(sub_n, "07:00-07:50")
        
        norm_allocs.append({
            "teacher_name": a.get("teacher_name") or a.get("teacher_id") or "Docente",
            "department": a.get("department") or "",
            "subject_name": a.get("subject_name") or a.get("subject") or "Disciplina",
            "shift": sh_code,
            "sub_slot": sub_n,
            "time_slot": slot_code,
            "time_interval": t_interval,
            "day_of_week": day_n
        })

    # Ordenar alocações por dia da semana (1-7), turno (M, T, N) e sub_slot (1-6)
    sorted_allocs = sorted(
        norm_allocs, 
        key=lambda x: (
            int(x["day_of_week"]), 
            shift_order.get(x["shift"], 1), 
            int(x["sub_slot"])
        )
    )

    # Calcular métricas rápidas (KPIs)
    total_aulas = len(sorted_allocs)
    unique_teachers = len(set(a.get("teacher_name", "Docente") for a in sorted_allocs)) if sorted_allocs else 0
    unique_shifts = sorted(list(set(str(a.get("shift", "M")).upper() for a in sorted_allocs)), key=lambda s: shift_order.get(s, 1))
    turnos_ocupados_str = ", ".join(shift_map.get(s, s) for s in unique_shifts) if unique_shifts else "Nenhum"
    acessivel_str = "Acessível (♿)" if is_accessible else "Não Acessível"

    if not HAS_REPORTLAB:
        lines = [
            "=" * 65,
            "CLASSSYNC AI - RELATÓRIO DE USO E OCUPAÇÃO DA SALA",
            "=" * 65,
            f"Sala: {room_name} ({block_id})",
            f"Tipo: {room_type} | Capacidade: {capacity} lugares | Acessibilidade: {acessivel_str}",
            f"Emissão: {datetime.datetime.now().strftime('%d/%m/%Y %H:%M')}",
            "-" * 65,
            f"MÉTRICAS: Total de Aulas: {total_aulas} | Docentes: {unique_teachers} | Turnos: {turnos_ocupados_str}",
            "-" * 65,
            "DOCENTES, DISCIPLINAS E HORÁRIOS:",
        ]
        for a in sorted_allocs:
            s_name = shift_map.get(str(a.get('shift', 'M')).upper(), a.get('shift', 'M'))
            d_name = days_map.get(int(a.get('day_of_week', 1)), f"Dia {a.get('day_of_week')}")
            sub_num = f"{int(a.get('sub_slot', 1)):02d}"
            lines.append(f"- Docente: {a.get('teacher_name')} ({a.get('department', '')}) | Disciplina: {a.get('subject_name')} | Dia: {d_name} | Turno: {s_name} | Horário: {a.get('time_interval')} (Slot: {a.get('time_slot')} - Aula {sub_num})")
        return "\n".join(lines).encode("utf-8")

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        rightMargin=30,
        leftMargin=30,
        topMargin=30,
        bottomMargin=30
    )

    styles = getSampleStyleSheet()
    
    app_header_style = ParagraphStyle(
        'AppHeader',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=9.5,
        textColor=colors.HexColor('#6366f1'),
        spaceAfter=3,
        textTransform='uppercase'
    )
    
    title_style = ParagraphStyle(
        'RoomMainTitle',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=16,
        textColor=colors.HexColor('#0f172a'),
        spaceAfter=4
    )
    
    meta_style = ParagraphStyle(
        'RoomMetaSub',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9,
        textColor=colors.HexColor('#475569'),
        spaceAfter=10
    )

    kpi_label_style = ParagraphStyle(
        'KPILabel',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=7.5,
        textColor=colors.HexColor('#64748b'),
        textTransform='uppercase'
    )
    kpi_val_blue = ParagraphStyle(
        'KPIValBlue',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=11,
        textColor=colors.HexColor('#0284c7')
    )
    kpi_val_green = ParagraphStyle(
        'KPIValGreen',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=11,
        textColor=colors.HexColor('#16a34a')
    )
    kpi_val_purple = ParagraphStyle(
        'KPIValPurple',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=11,
        textColor=colors.HexColor('#9333ea')
    )

    footer_style = ParagraphStyle(
        'DocFooter',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=7.5,
        textColor=colors.HexColor('#94a3b8'),
        alignment=1,
        spaceBefore=12
    )

    elements = []

    # 1. Cabeçalho Principal
    elements.append(Paragraph("ClassSync AI — Gestão de Espaços e Alocação", app_header_style))
    elements.append(Paragraph(f"{room_name} ({block_id})", title_style))
    elements.append(Paragraph(
        f"<b>Tipo:</b> {room_type} &nbsp;|&nbsp; <b>Capacidade:</b> {capacity} lugares &nbsp;|&nbsp; <b>Acessibilidade:</b> <font color='#16a34a'><b>{acessivel_str}</b></font> &nbsp;|&nbsp; <b>Emissão:</b> {datetime.datetime.now().strftime('%d/%m/%Y às %H:%M')}",
        meta_style
    ))
    elements.append(Spacer(1, 2))

    # 2. Quadro de Métricas / Resumo Rápido (KPI Cards)
    kpi_col1 = [
        Paragraph("TOTAL DE AULAS", kpi_label_style),
        Paragraph(f"{total_aulas} {'aula alocada' if total_aulas == 1 else 'aulas alocadas'}", kpi_val_blue)
    ]
    kpi_col2 = [
        Paragraph("DOCENTES DISTINTOS", kpi_label_style),
        Paragraph(f"{unique_teachers}", kpi_val_green)
    ]
    kpi_col3 = [
        Paragraph("TURNOS OCUPADOS", kpi_label_style),
        Paragraph(turnos_ocupados_str, kpi_val_purple)
    ]

    kpi_table = Table([[kpi_col1, kpi_col2, kpi_col3]], colWidths=[184, 184, 184])
    kpi_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#f8fafc')),
        ('BOX', (0, 0), (-1, -1), 1, colors.HexColor('#e2e8f0')),
        ('INNERGRID', (0, 0), (-1, -1), 1, colors.HexColor('#e2e8f0')),
        ('PADDING', (0, 0), (-1, -1), 7),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
    ]))
    elements.append(kpi_table)
    elements.append(Spacer(1, 12))

    # 3. Tabela de Detalhamento com todas as colunas da imagem
    table_data = [[
        Paragraph("<font color='white'><b>Docente</b></font>", styles['Normal']),
        Paragraph("<font color='white'><b>Disciplina</b></font>", styles['Normal']),
        Paragraph("<font color='white'><b>Dia</b></font>", styles['Normal']),
        Paragraph("<font color='white'><b>Turno</b></font>", styles['Normal']),
        Paragraph("<font color='white'><b>Horário / Slot</b></font>", styles['Normal'])
    ]]

    shift_colors_map = {
        "Manhã": ('#eff6ff', '#1d4ed8'),
        "Tarde": ('#fffbeb', '#b45309'),
        "Noite": ('#faf5ff', '#7e22ce')
    }

    if sorted_allocs:
        for a in sorted_allocs:
            t_name = a.get('teacher_name', 'Docente')
            t_dept = a.get('department', '')
            docente_cell = f"<b>{t_name}</b>"
            if t_dept:
                docente_cell += f"<br/><font size=7.5 color='#64748b'>{t_dept}</font>"

            subj_name = a.get('subject_name', 'Disciplina Geral')
            disciplina_cell = f"<b>{subj_name}</b>"

            day_name = days_map.get(int(a.get("day_of_week", 1)), f"Dia {a.get('day_of_week')}")

            shift_code = str(a.get("shift", "M")).upper()
            shift_name = shift_map.get(shift_code, shift_code)
            shift_bg, shift_fg = shift_colors_map.get(shift_name, ('#eff6ff', '#1d4ed8'))
            turno_cell = f"<font color='{shift_fg}'><b>{shift_name}</b></font>"

            time_interval = a.get("time_interval", "07:00-07:50")
            sub_num = f"{int(a.get('sub_slot', 1)):02d}"
            slot_code = a.get("time_slot", f"{shift_code}{a.get('sub_slot', 1)}")
            horario_cell = f"<b>{time_interval}</b><br/><font size=7.5 color='#64748b'>Slot: {slot_code} (Aula {sub_num})</font>"

            table_data.append([
                Paragraph(docente_cell, styles['Normal']),
                Paragraph(disciplina_cell, styles['Normal']),
                Paragraph(day_name, styles['Normal']),
                Paragraph(turno_cell, styles['Normal']),
                Paragraph(horario_cell, styles['Normal'])
            ])
    else:
        table_data.append([
            Paragraph("<i>Nenhuma aula ou ocupação registrada para este espaço no momento.</i>", styles['Normal']),
            "-", "-", "-", "-"
        ])

    data_table = Table(table_data, colWidths=[140, 145, 87, 65, 115])
    data_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#0f172a')),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 8),
        ('TOPPADDING', (0, 0), (-1, 0), 8),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f8fafc')]),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#cbd5e1')),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('PADDING', (0, 0), (-1, -1), 6),
    ]))
    elements.append(data_table)

    # 4. Rodapé institucional
    elements.append(Spacer(1, 15))
    elements.append(Paragraph(
        "Documento gerado automaticamente pelo <b>ClassSync AI</b> • Relatório Oficial de Auditoria e Ocupação Predial",
        footer_style
    ))

    doc.build(elements)
    pdf_bytes = buffer.getvalue()
    buffer.close()
    return pdf_bytes
